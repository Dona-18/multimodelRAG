"""Tests for lesson-style LangChain RAG; only embeddings use local Ollama."""
from __future__ import annotations

from pathlib import Path

from docx import Document as DocxDocument
from langchain_core.messages import AIMessage
from langchain_core.runnables import RunnableLambda
from pypdf import PdfWriter

from src.generation import INSUFFICIENT_EVIDENCE_MESSAGE, answer_question, stream_answer
from src.ingestion import load_document, make_chunks, recursive_character_chunks
from src.retrieval import generate_query_variations, hybrid_search, reciprocal_rank_fusion
from src.state import (
    chat_is_unlocked,
    clear_active_source,
    load_active_source,
    load_index_settings,
    requires_reingestion,
    save_active_source,
    save_index_settings,
)
from src.types import (
    Chunk,
    GenerationSettings,
    IndexSettings,
    RetrievalResult,
    RetrievalSettings,
    SearchHit,
)
from src.vector_store import VectorStore


def chunk(identifier: str, text: str, page: int | None = 1) -> Chunk:
    return Chunk(identifier, text, "notes.txt", "source-1", page, 0)


class FakeChatModel:
    """Small ChatOllama-shaped double for generation-only unit tests."""

    def invoke(self, messages):
        assert "[notes.txt, p. 1]" in str(messages[-1].content)
        return AIMessage(content="Neural networks learn from examples. [notes.txt, p. 1]")

    def stream(self, messages):
        assert "[notes.txt, p. 1]" in str(messages[-1].content)
        yield AIMessage(content="Neural ")
        yield AIMessage(content="networks learn. [notes.txt, p. 1]")


class InMemoryStore:
    """Deterministic retrieval double for checking multi-query RRF behavior."""

    def __init__(self, chunks: list[Chunk]) -> None:
        self.chunks = chunks

    def all_chunks(self) -> list[Chunk]:
        return self.chunks

    def vector_search(self, query: str, k: int) -> list[Chunk]:
        return self.chunks if "alternate" in query else list(reversed(self.chunks))

    def vector_similarity_scores(self, query: str, k: int) -> dict[str, float]:
        return {self.chunks[0].id: 0.9, self.chunks[1].id: 0.4}


def test_text_docx_and_pdf_loading_preserves_metadata(tmp_path: Path) -> None:
    text_path = tmp_path / "notes.txt"
    text_path.write_text("A short local note about retrieval.")
    text_pages = load_document(text_path)
    assert text_pages[0].source_name == "notes.txt"
    assert text_pages[0].page is None

    docx_path = tmp_path / "table.docx"
    docx = DocxDocument()
    docx.add_paragraph("The project budget is 42 dollars.")
    table = docx.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Owner"
    table.cell(0, 1).text = "Ari"
    docx.save(docx_path)
    assert "Owner | Ari" in load_document(docx_path)[0].text

    pdf_path = tmp_path / "empty.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    writer.add_blank_page(width=72, height=72)
    with pdf_path.open("wb") as output:
        writer.write(output)
    pdf_pages = load_document(pdf_path)
    assert [page.page for page in pdf_pages] == [1, 2]
    assert pdf_pages[0].source_name == "empty.pdf"


def test_recursive_chunking_overlap_and_page_reference(tmp_path: Path) -> None:
    text = "First paragraph has several useful words.\n\nSecond paragraph has distinct retrieval evidence."
    pieces = recursive_character_chunks(text, chunk_size=55, overlap=10)
    assert len(pieces) >= 2
    assert all(len(piece) <= 55 for piece in pieces)

    path = tmp_path / "page.txt"
    path.write_text(text)
    chunks = make_chunks(load_document(path), 55, 10)
    assert len(chunks) == len(pieces)
    assert all(item.page is None for item in chunks)
    assert all(item.chroma_metadata()["page_available"] is False for item in chunks)


def test_rrf_uses_exact_lesson_formula_and_keeps_duplicate_text_sources() -> None:
    first = chunk("first", "Same sentence", page=1)
    second = Chunk("second", "Same sentence", "notes.txt", "source-1", 2, 1)
    fused = reciprocal_rank_fusion([[(first, 0.9), (second, 0.8)], [(second, 0.9), (first, 0.8)]])
    assert len(fused) == 2
    assert fused[0][1] == (1 / 61) + (1 / 62)


def test_chroma_vector_bm25_rrf_pipeline(tmp_path: Path) -> None:
    store = VectorStore(tmp_path / "chroma", "unit_test")
    neural = chunk("neural", "Neural networks use layers and training examples.")
    plants = Chunk("plants", "Photosynthesis lets plants make food from light.", "biology.txt", "source-2", 2, 1)
    # This is a real compatibility check: LangChain OllamaEmbeddings → Chroma.
    assert store.add_chunks([neural, plants]) == 2
    assert store.count == 2

    result = hybrid_search("How do neural networks learn?", store, top_k=2, min_similarity=0.0)
    assert result.sufficient
    assert result.hits[0].chunk.id == "neural"
    assert result.hits[0].bm25_score is not None
    assert result.hits[0].vector_similarity is not None


def test_multi_query_variations_and_rrf_are_available_without_extra_models() -> None:
    json_model = RunnableLambda(lambda _: '{"queries": ["alternate neural learning question"]}')
    queries = generate_query_variations("How do neural networks learn?", llm=json_model, count=1)
    assert queries == ["How do neural networks learn?", "alternate neural learning question"]

    neural = chunk("neural", "Neural networks learn from examples.")
    plants = Chunk("plants", "Plants use photosynthesis.", "biology.txt", "source-2", 2, 1)
    result = hybrid_search(
        "How do neural networks learn?",
        InMemoryStore([neural, plants]),
        top_k=2,
        min_similarity=0.1,
        use_multi_query=True,
        llm=json_model,
    )
    assert result.sufficient
    assert "multi-query hybrid RRF" in result.reason


def test_generation_refuses_before_model_for_weak_evidence() -> None:
    retrieval = RetrievalResult(sufficient=False, reason="Below threshold")
    result = answer_question("What is unknown?", retrieval)
    assert result.answer == INSUFFICIENT_EVIDENCE_MESSAGE
    assert not result.citations


def test_generation_includes_page_citation_in_grounded_prompt() -> None:
    hit = SearchHit(chunk("neural", "Neural networks learn from examples."), 0.9, 1.0, 0.03)
    result = answer_question("How do networks learn?", RetrievalResult([hit], True), llm=FakeChatModel())
    assert "[notes.txt, p. 1]" in result.answer
    assert result.citations == ["notes.txt, p. 1"]


def test_chat_stays_locked_until_ingestion_creates_evidence() -> None:
    assert not chat_is_unlocked(0)
    assert chat_is_unlocked(1)


def test_advanced_settings_validate_and_apply_to_streaming_path() -> None:
    generation = GenerationSettings(temperature=0.2, top_p=0.9, max_output_tokens=256, context_length=4096, streaming=True)
    retrieval = RetrievalSettings(top_k=5, use_multi_query=True, use_hybrid_search=False)
    index = IndexSettings(chunk_size=1000, chunk_overlap=150)
    generation.validate()
    retrieval.validate()
    index.validate()

    hit = SearchHit(chunk("neural", "Neural networks learn from examples."), 0.9, 1.0, 0.03)
    pieces = list(stream_answer("How do networks learn?", RetrievalResult([hit], True), generation, llm=FakeChatModel()))
    assert "".join(pieces) == "Neural networks learn. [notes.txt, p. 1]"


def test_chunk_settings_change_requires_reingestion(monkeypatch, tmp_path: Path) -> None:
    state_path = tmp_path / "index_settings.json"
    monkeypatch.setattr("src.state.INDEX_SETTINGS_PATH", state_path)
    original = IndexSettings(chunk_size=1000, chunk_overlap=150)
    save_index_settings(original)
    assert load_index_settings() == original
    assert not requires_reingestion(2, original)
    assert requires_reingestion(2, IndexSettings(chunk_size=800, chunk_overlap=100))


def test_active_source_keeps_only_the_latest_project_upload(monkeypatch, tmp_path: Path) -> None:
    uploads = tmp_path / "uploads"
    uploads.mkdir()
    state_path = tmp_path / "active_source.json"
    monkeypatch.setattr("src.state.UPLOADS_DIR", uploads)
    monkeypatch.setattr("src.state.ACTIVE_SOURCE_PATH", state_path)

    first = uploads / "first.txt"
    second = uploads / "second.txt"
    first.write_text("First document")
    second.write_text("Replacement document")
    save_active_source(first)
    assert load_active_source() == first.resolve()

    save_active_source(second)
    assert load_active_source() == second.resolve()
    clear_active_source()
    assert load_active_source() is None
