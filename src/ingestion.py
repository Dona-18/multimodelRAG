"""Document ingestion: load pasted text, text files, docx, PDFs, and images."""
import io
import json
from pathlib import Path
from typing import List, Dict, Any, Union

from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pypdf import PdfReader
from PIL import Image
import ollama

from src.config import LLM_MODEL, OLLAMA_URL


def describe_image(image_bytes: bytes) -> str:
    """Analyze image, chart, or diagram with local vision LLM."""
    try:
        pil_img = Image.open(io.BytesIO(image_bytes)).convert("RGB")
        # Skip tiny decorative icons (< 60px)
        if pil_img.width < 60 or pil_img.height < 60:
            return ""
        buf = io.BytesIO()
        pil_img.save(buf, format="JPEG")
        jpeg_bytes = buf.getvalue()

        client = ollama.Client(host=OLLAMA_URL)
        response = client.chat(
            model=LLM_MODEL,
            messages=[{
                "role": "user",
                "content": (
                    "Describe this document image, chart, or diagram concisely for retrieval. "
                    "Include any visible headings, labels, numbers, and relationships."
                ),
                "images": [jpeg_bytes],
            }],
        )
        return response["message"]["content"].strip()
    except Exception:
        return ""


def load_pasted_text(text: str, source_name: str = "Pasted_Note.txt") -> List[Document]:
    """Turn pasted text into a LangChain Document."""
    cleaned = text.strip()
    if not cleaned:
        return []
    return [
        Document(
            page_content=cleaned,
            metadata={
                "source": source_name,
                "page": 1,
                "type": "text",
                "original_content": json.dumps({"raw_text": cleaned}),
            },
        )
    ]


def load_text_file(file_path: Union[str, Path]) -> List[Document]:
    """Read a plain text or markdown file."""
    path = Path(file_path)
    text = path.read_text(encoding="utf-8", errors="ignore")
    return [
        Document(
            page_content=text,
            metadata={
                "source": path.name,
                "page": 1,
                "type": "text",
                "original_content": json.dumps({"raw_text": text}),
            },
        )
    ]


def load_docx_file(file_path: Union[str, Path]) -> List[Document]:
    """Read a docx document."""
    try:
        from docx import Document as DocxDocument
        doc = DocxDocument(str(file_path))
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        text = "\n\n".join(paragraphs)
    except Exception:
        text = Path(file_path).read_text(encoding="utf-8", errors="ignore")

    path = Path(file_path)
    return [
        Document(
            page_content=text,
            metadata={
                "source": path.name,
                "page": 1,
                "type": "text",
                "original_content": json.dumps({"raw_text": text}),
            },
        )
    ]


def load_image_file(file_path: Union[str, Path]) -> List[Document]:
    """Read an image file (PNG, JPG, etc.) and describe it using the vision model."""
    path = Path(file_path)
    image_bytes = path.read_bytes()
    description = describe_image(image_bytes)
    content = description if description else f"Image file: {path.name}"
    return [
        Document(
            page_content=f"[Image {path.name}]: {content}",
            metadata={
                "source": path.name,
                "page": 1,
                "type": "image",
                "original_content": json.dumps({"visual_description": content}),
            },
        )
    ]


def load_pdf_file(file_path: Union[str, Path], include_vision: bool = True) -> List[Document]:
    """Extract text and visual diagrams from a PDF file page by page."""
    path = Path(file_path)
    reader = PdfReader(str(path))
    docs: List[Document] = []

    for page_idx, page in enumerate(reader.pages, start=1):
        # 1. Extract text
        text = (page.extract_text() or "").strip()
        if text:
            docs.append(
                Document(
                    page_content=text,
                    metadata={
                        "source": path.name,
                        "page": page_idx,
                        "type": "text",
                        "original_content": json.dumps({"raw_text": text}),
                    },
                )
            )

        # 2. Extract and describe diagrams / images on this page
        if include_vision and hasattr(page, "images"):
            for img in page.images:
                if len(img.data) < 1024:  # skip tiny decorative icons
                    continue
                desc = describe_image(img.data)
                if desc:
                    docs.append(
                        Document(
                            page_content=f"[Visual / Diagram on Page {page_idx}]: {desc}",
                            metadata={
                                "source": path.name,
                                "page": page_idx,
                                "type": "image",
                                "original_content": json.dumps({"visual_description": desc}),
                            },
                        )
                    )
    return docs


def load_any_document(file_path: Union[str, Path]) -> List[Document]:
    """Route document based on file extension."""
    path = Path(file_path)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return load_pdf_file(path)
    elif suffix == ".docx":
        return load_docx_file(path)
    elif suffix in {".png", ".jpg", ".jpeg", ".webp"}:
        return load_image_file(path)
    return load_text_file(path)


def split_documents(
    documents: List[Document],
    chunk_size: int = 1000,
    chunk_overlap: int = 150,
) -> List[Document]:
    """Split documents into chunks using RecursiveCharacterTextSplitter."""
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
    )
    return splitter.split_documents(documents)


# Multimodal helpers matching 8_multi_modal_rag.ipynb
def partition_document(file_path: Union[str, Path], strategy: str = "auto") -> List[Any]:
    """Extract elements from PDF (supports unstructured if installed, or fallback to pypdf)."""
    try:
        from unstructured.partition.pdf import partition_pdf
        return partition_pdf(filename=str(file_path), strategy=strategy)
    except ImportError:
        return load_pdf_file(file_path)


def create_chunks_by_title(elements: List[Any]) -> List[Any]:
    """Chunk elements by title if unstructured is available, or return as-is."""
    try:
        from unstructured.chunking.title import chunk_by_title
        return chunk_by_title(elements, max_characters=1500, combine_text_under_n_chars=200)
    except ImportError:
        return elements


def separate_content_types(chunk: Any) -> Dict[str, Any]:
    """Inspect content types in a chunk (text, tables, images)."""
    text = getattr(chunk, "text", getattr(chunk, "page_content", str(chunk)))
    return {"text": text, "tables": [], "images": [], "types": ["text"]}


def summarise_chunks(chunks: List[Any], source_name: str = "doc") -> List[Document]:
    """Convert chunks into LangChain Documents."""
    docs: List[Document] = []
    for chunk in chunks:
        if isinstance(chunk, Document):
            docs.append(chunk)
            continue
        text = getattr(chunk, "text", str(chunk))
        page = getattr(getattr(chunk, "metadata", None), "page_number", 1) or 1
        docs.append(
            Document(
                page_content=text,
                metadata={
                    "source": source_name,
                    "page": page,
                    "type": "text",
                    "original_content": json.dumps({"raw_text": text}),
                },
            )
        )
    return docs
